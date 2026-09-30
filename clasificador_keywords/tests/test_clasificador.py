import asyncio
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from clasificador.entrada import Negocio, leer_keywords
from clasificador.excel import exportar
from clasificador.jev import clasificar, preguntas
from clasificador.prioridad import priorizar

RAIZ = Path(__file__).resolve().parents[1]
NEGOCIO = Negocio.cargar(RAIZ / "negocios/sprint_latam.yaml")


class ClienteFalso:
    """Imita AsyncTypeSafeClient con respuestas fijas según palabras de la keyword."""

    def __init__(self):
        self.llamadas = 0

    async def system_one(self, state, questions):
        self.llamadas += 1
        kw = state["keyword"]
        assert set(questions) == {"intencion", "etapa", "servicio", "relevancia", "competidor", "no_cliente", "otro_pais"}
        if "falla" in kw:
            raise RuntimeError("error de red")
        curso = "curso" in kw
        agencia = "agencia" in kw
        eleccion = lambda v, conf=0.9: SimpleNamespace(choice=v, confidence=conf)
        return SimpleNamespace(
            model="jev-test",
            usage=SimpleNamespace(input_tokens=1000, output_tokens=50),
            choices={
                "intencion": eleccion("transaccional" if agencia else "informacional"),
                "etapa": eleccion("bofu" if agencia else "tofu"),
                "servicio": eleccion("seo", 0.3 if "dudosa" in kw else 0.9),
            },
            scores={"relevancia": SimpleNamespace(score=3.0 if agencia else 2.0, confidence=0.9)},
            nouls={
                "competidor": SimpleNamespace(noul=0.9 if "onza" in kw else 0.05),
                "no_cliente": SimpleNamespace(noul=0.95 if curso else 0.05),
                "otro_pais": SimpleNamespace(noul=0.9 if "madrid" in kw else 0.05),
            },
        )


def test_lee_csv_de_ubersuggest_con_columnas_distintas(tmp_path):
    csv = tmp_path / "kw.csv"
    csv.write_text("Keyword;Volume;SD;CPC\nagencia seo;1.600;13;2,5\nAgencia SEO;10;1;0\nseo;90;;\n", encoding="utf-8")
    kws = leer_keywords(csv)
    assert [k.keyword for k in kws] == ["agencia seo", "seo"]  # sin duplicados
    assert kws[0].volumen == 1600 and kws[0].dificultad == 13 and kws[0].cpc == 2.5
    assert kws[1].dificultad is None


def test_ejemplo_de_sprint_se_lee():
    kws = leer_keywords(RAIZ / "ejemplos/sprint_keywords_chile.csv")
    assert len(kws) > 100
    assert any(k.keyword == "agencia seo chile" and k.volumen == 260 for k in kws)


def test_preguntas_incluyen_servicios_del_negocio():
    servicio = preguntas(NEGOCIO)["servicio"]
    assert "geo" in servicio.criteria and "ninguno" in servicio.criteria


def test_clasifica_prioriza_y_calcula_costo():
    from clasificador.entrada import Keyword

    kws = [
        Keyword("agencia seo chile", 260, 13),
        Keyword("que es posicionamiento web", 10, 5),
        Keyword("curso google ads", 210, 21),
        Keyword("onza agencia de marketing digital", 110, 7),
        Keyword("agencia seo madrid", 90, 21),
        Keyword("agencia dudosa", 50, 10),
        Keyword("falla", 10, 10),
    ]
    cliente = ClienteFalso()
    resultado, costo = asyncio.run(clasificar(cliente, NEGOCIO, kws))
    resultado = {c.keyword.keyword: c for c in priorizar(resultado)}

    assert resultado["agencia seo chile"].decision == "atacar"
    assert "servicio" in resultado["agencia seo chile"].motivo
    assert resultado["que es posicionamiento web"].decision == "atacar"
    assert resultado["agencia seo chile"].prioridad > resultado["que es posicionamiento web"].prioridad
    assert resultado["curso google ads"].decision == "excluir"
    assert resultado["onza agencia de marketing digital"].decision == "excluir"
    assert resultado["agencia seo madrid"].decision == "excluir"
    assert resultado["agencia dudosa"].decision == "revisar"
    assert resultado["falla"].decision == "sin clasificar"

    assert costo.solicitudes == 7 and costo.fallidas == 1
    assert costo.tokens_entrada == 6000 and costo.modelo == "jev-test"
    assert abs(costo.usd - 6000 * 0.042 / 1e6) < 1e-12


def test_presupuesto_detiene_consultas():
    from clasificador.entrada import Keyword

    kws = [Keyword(f"agencia {i}", 10, 10) for i in range(30)]
    cliente = ClienteFalso()
    # cada consulta cuesta 1000 tokens = US$0.000042; con este tope caben unas pocas
    _, costo = asyncio.run(clasificar(cliente, NEGOCIO, kws, presupuesto_usd=0.0001))
    assert costo.omitidas_por_presupuesto > 0
    assert cliente.llamadas + costo.omitidas_por_presupuesto == 30


def test_exporta_excel(tmp_path):
    kws = leer_keywords(RAIZ / "ejemplos/sprint_keywords_chile.csv")[:12]
    resultado, costo = asyncio.run(clasificar(ClienteFalso(), NEGOCIO, kws))
    destino = exportar(priorizar(resultado), costo, tmp_path / "out.xlsx", NEGOCIO.nombre)
    libro = load_workbook(destino)
    assert libro.sheetnames == ["Keywords", "Prioridades", "Servicio x etapa", "Costo"]
    assert libro["Keywords"].max_row == 13
    assert libro["Costo"]["B10"].value == round(costo.usd, 6)
