"""Exporta la clasificación a un Excel con varias hojas."""

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .jev import USD_POR_MILLON_ENTRADA, Clasificacion, Costo

COLOR_DECISION = {
    "atacar": "C6EFCE",
    "revisar": "FFEB9C",
    "baja relevancia": "EDEDED",
    "excluir": "FFC7CE",
    "sin clasificar": "FFC7CE",
}
ENCABEZADO = PatternFill("solid", fgColor="040A14")


def _hoja(ws, encabezados: list[str], filas: list[list], anchos: list[int] | None = None) -> None:
    ws.append(encabezados)
    for celda in ws[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = ENCABEZADO
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for fila in filas:
        ws.append(fila)
    ws.freeze_panes = "A2"
    if filas:
        ws.auto_filter.ref = ws.dimensions
    for i, ancho in enumerate(anchos or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _pct(x: float) -> float:
    return round(x * 100)


def exportar(clasificaciones: list[Clasificacion], costo: Costo, destino: str | Path, negocio: str) -> Path:
    destino = Path(destino)
    libro = Workbook()

    # 1. Todas las keywords
    ws = libro.active
    ws.title = "Keywords"
    filas = []
    for c in clasificaciones:
        kw = c.keyword
        filas.append([
            kw.keyword, kw.volumen, kw.dificultad, kw.cpc, c.decision, c.prioridad, c.motivo,
            c.intencion, _pct(c.intencion_confianza), c.etapa, c.servicio, _pct(c.servicio_confianza),
            round(c.relevancia, 2), _pct(c.competidor), _pct(c.no_cliente), _pct(c.otro_pais),
        ])
    _hoja(
        ws,
        ["Keyword", "Volumen", "Dificultad", "CPC USD", "Decisión", "Prioridad", "Motivo / acción",
         "Intención", "Conf. intención %", "Etapa", "Servicio", "Conf. servicio %",
         "Relevancia (0-3)", "Otra marca %", "No cliente %", "Otro país %"],
        filas,
        [38, 10, 10, 9, 15, 10, 42, 14, 10, 8, 18, 10, 11, 10, 10, 10],
    )
    for fila in ws.iter_rows(min_row=2):
        color = COLOR_DECISION.get(fila[4].value)
        if color:
            fila[4].fill = PatternFill("solid", fgColor=color)

    # 2. Solo las que conviene atacar, en orden
    ws = libro.create_sheet("Prioridades")
    atacar = [c for c in clasificaciones if c.decision == "atacar"]
    _hoja(
        ws,
        ["#", "Keyword", "Volumen", "Dificultad", "Prioridad", "Servicio", "Etapa", "Qué crear"],
        [[i, c.keyword.keyword, c.keyword.volumen, c.keyword.dificultad, c.prioridad, c.servicio, c.etapa, c.motivo]
         for i, c in enumerate(atacar, start=1)],
        [5, 38, 10, 10, 10, 18, 8, 42],
    )

    # 3. Mapa servicio x etapa (keywords a atacar y volumen)
    ws = libro.create_sheet("Servicio x etapa")
    mapa = defaultdict(lambda: [0, 0])
    for c in atacar:
        mapa[(c.servicio, c.etapa)][0] += 1
        mapa[(c.servicio, c.etapa)][1] += c.keyword.volumen
    servicios = sorted({s for s, _ in mapa})
    filas = []
    for s in servicios:
        fila = [s]
        for etapa in ("tofu", "mofu", "bofu"):
            n, vol = mapa.get((s, etapa), (0, 0))
            fila += [n, vol]
        filas.append(fila)
    _hoja(
        ws,
        ["Servicio", "TOFU kws", "TOFU vol.", "MOFU kws", "MOFU vol.", "BOFU kws", "BOFU vol."],
        filas,
        [20, 10, 10, 10, 10, 10, 10],
    )

    # 4. Costo de la ejecución
    ws = libro.create_sheet("Costo")
    _hoja(
        ws,
        ["Concepto", "Valor"],
        [
            ["Negocio", negocio],
            ["Modelo Jev", costo.modelo or "jev-latest"],
            ["Keywords clasificadas", costo.solicitudes - costo.fallidas],
            ["Solicitudes fallidas", costo.fallidas],
            ["Omitidas por presupuesto", costo.omitidas_por_presupuesto],
            ["Tokens de entrada", costo.tokens_entrada],
            ["Tokens de salida (no se cobran)", costo.tokens_salida],
            ["Precio USD por millón de tokens de entrada", USD_POR_MILLON_ENTRADA],
            ["Costo total USD", round(costo.usd, 6)],
        ] + [["Error", e] for e in costo.errores[:50]],
        [42, 30],
    )

    destino.parent.mkdir(parents=True, exist_ok=True)
    libro.save(destino)
    return destino
